"""Extract paragraphs and table cells from Word documents."""

from pathlib import Path

from docx import Document


def extract(path: Path, max_chars: int) -> str:
    document = Document(str(path))
    chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        chunks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(chunks)[:max_chars]
