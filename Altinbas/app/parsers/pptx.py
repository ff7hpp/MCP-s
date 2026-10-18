"""Extract text and table cells from PowerPoint slides."""

from pathlib import Path

from pptx import Presentation


def extract(path: Path, max_chars: int) -> str:
    presentation = Presentation(str(path))
    chunks: list[str] = []
    total = 0
    for number, slide in enumerate(presentation.slides, 1):
        pieces = [f"Slide {number}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                pieces.extend(
                    paragraph.text
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                )
            if shape.has_table:
                pieces.extend(
                    " | ".join(cell.text for cell in row.cells) for row in shape.table.rows
                )
        chunk = "\n".join(pieces)
        chunks.append(chunk[: max_chars - total])
        total += len(chunks[-1])
        if total >= max_chars:
            break
    return "\n\n".join(chunks)
