from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from app.parsers import docx, pptx, text
from app.services.documents import read_document


def test_text_extractors(tmp_path: Path) -> None:
    plain = tmp_path / "note.md"
    plain.write_text("Lecture notes", encoding="utf-8")
    assert text.extract(plain, 100) == "Lecture notes"

    word = Document()
    word.add_paragraph("Validation topic")
    word_path = tmp_path / "note.docx"
    word.save(word_path)
    assert "Validation topic" in docx.extract(word_path, 100)

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 1000000, 1000000)
    box.text = "Software testing"
    slides_path = tmp_path / "lecture.pptx"
    slides.save(slides_path)
    assert "Software testing" in pptx.extract(slides_path, 100)


def test_text_limit_rejected() -> None:
    with pytest.raises(ValueError, match="TEXT_LIMIT_INVALID"):
        read_document("SWE431", "Lecture1.pdf", 0)
