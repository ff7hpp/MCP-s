from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter

from app.config import Config
from app.parsers import docx, pptx, text
from app.services import documents
from app.services.documents import read_document


def test_text_extractors(tmp_path: Path) -> None:
    plain = tmp_path / "note.md"
    plain.write_text("Lecture notes", encoding="utf-8")
    assert text.extract(plain, 100) == "Lecture notes"

    word = Document()
    word.add_paragraph("Validation topic")
    word_path = tmp_path / "note.docx"
    word.save(word_path)
    assert "Validation topic" in docx.extract(word_path, 100)

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 1000000, 1000000)
    box.text = "Software testing"
    slides_path = tmp_path / "lecture.pptx"
    slides.save(slides_path)
    assert "Software testing" in pptx.extract(slides_path, 100)


def test_text_limit_rejected() -> None:
    with pytest.raises(ValueError, match="TEXT_LIMIT_INVALID"):
        read_document("SWE431", "Lecture1.pdf", 0)


def test_syllabus_evidence_weights_and_tracking_without_guesses(tmp_path, monkeypatch):
    course = tmp_path / "University" / "SWE431 Testing"
    course.mkdir(parents=True)
    syllabus = course / "syllabus.txt"
    syllabus.write_text(
        "Assessment\nMidterm 30%\nFinal Exam 70%\nAssignments: two reports\nSubmit through LMS. No late submissions.\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(
        documents, "load_config", lambda: Config(term="Fall 2026/2027", drive_root=tmp_path)
    )
    path = tmp_path / "syllabus-state.json"
    first = documents.inspect_syllabus("SWE431", "syllabus.txt", path)
    assert first["change_status"] == "BASELINE"
    assert [row["percent"] for row in first["grade_weights"]] == [30, 70]
    assert (
        first["evidence"]["submission_rules"][0]["quote"]
        == "Submit through LMS. No late submissions."
    )
    assert (
        documents.inspect_syllabus("SWE431", "syllabus.txt", path)["change_status"] == "UNCHANGED"
    )
    syllabus.write_text("The grading scheme is not announced.\n", encoding="utf-8")
    changed = documents.inspect_syllabus("SWE431", "syllabus.txt", path)
    assert changed["change_status"] == "UPDATED"
    assert changed["grade_weights"] == "Unknown"
    assert changed["assignments"] == "Unknown"


def test_material_search_deduplicates_and_reports_no_text_and_bad_pdf(tmp_path, monkeypatch):
    course = tmp_path / "University" / "SWE431 Testing"
    course.mkdir(parents=True)
    (course / "notes.txt").write_text("Boundary value analysis and testing", encoding="utf-8")
    (course / "duplicate.txt").write_text("Boundary value analysis and testing", encoding="utf-8")
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(course / "image.pdf")
    (course / "bad.pdf").write_bytes(b"%PDF-invalid")
    monkeypatch.setattr(
        documents, "load_config", lambda: Config(term="Fall 2026/2027", drive_root=tmp_path)
    )
    result = documents.search_materials("SWE431", "boundary")
    assert len(result["matches"]) == 1
    assert result["matches"][0]["source"].startswith(str(course))
    assert result["completeness"] == "PARTIAL"
    assert len(result["errors"]) == 2


def test_receipt_from_other_drive_root_does_not_block_selected_account(tmp_path, monkeypatch):
    from app.drive import records
    from app.storage.state import save_state

    root = tmp_path / "chosen"
    course = root / "University" / "SWE431 Testing"
    course.mkdir(parents=True)
    selected = course / "lecture.pdf"
    selected.write_bytes(b"selected file")
    outside = tmp_path / "other-account" / "lecture.pdf"
    save_state(
        records.STATE_PATH,
        {
            "version": 1,
            "sections": {
                "old": {
                    "course_code": "SWE431",
                    "filename": "lecture.pdf",
                    "term": "Fall 2026/2027",
                    "path": str(outside),
                    "revision": "legacy",
                    "sha256": "0" * 64,
                    "size": 1,
                }
            },
        },
    )
    monkeypatch.setattr(
        documents, "load_config", lambda: Config(term="Fall 2026/2027", drive_root=root)
    )
    assert documents.find_file("SWE431", "lecture.pdf") == selected
